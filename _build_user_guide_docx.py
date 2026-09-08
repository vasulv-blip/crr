#!/usr/bin/env python3
"""Build numbered step-by-step Word user guide for the CR simulation software."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT_APP = Path(r"C:\work\projectx\embdeeded\proposalprojects\cognitive_radio_sim") / "Cognitive_Radio_Simulation_User_Guide.docx"
OUT_DESK = Path.home() / "Desktop" / "proposals" / "Cognitive_Radio_Simulation_User_Guide.docx"

NAVY = RGBColor(0x0B, 0x2A, 0x5B)
GRAY = RGBColor(0x55, 0x55, 0x55)


def set_run(run, size=11, bold=False, italic=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color


def add_center(doc, text, size=11, bold=False, color=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run(text), size=size, bold=bold, color=color)


def add_p(doc, text, size=11, bold=False, italic=False):
    p = doc.add_paragraph()
    set_run(p.add_run(text), size=size, bold=bold, italic=italic)


def add_num(doc, text, size=11):
    p = doc.add_paragraph(text, style="List Number")
    for r in p.runs:
        set_run(r, size=size)


def add_bullet(doc, text, size=11):
    p = doc.add_paragraph(text, style="List Bullet")
    for r in p.runs:
        set_run(r, size=size)


def main():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.75)
        s.bottom_margin = Inches(0.75)
        s.left_margin = Inches(0.85)
        s.right_margin = Inches(0.85)

    add_center(doc, "eAge Innovations", 12, True, NAVY)
    add_center(doc, "Cognitive Radio Software Simulation", 14, True, NAVY)
    add_center(doc, "User Guide — Step-by-Step Operation", 12, True)
    add_center(doc, "How to start the software and walk through demonstration scenarios", 10, False, GRAY)
    add_center(doc, "September 2026", 10, False, GRAY)

    add_p(
        doc,
        "This guide explains how to run the Cognitive Radio simulation on a personal computer "
        "and how to operate it step by step during a technical demonstration. "
        "All steps are numbered in order.",
    )

    # 1
    doc.add_heading("1. Purpose of this software", level=1)
    add_num(doc, "Simulate a configurable number of radio channels on a PC.")
    add_num(doc, "Inject jam profiles and observe which channels become FREE, DEGRADED, or BLOCKED.")
    add_num(doc, "Show how the Cognitive Radio engine recommends a channel or advises increasing power.")
    add_num(doc, "Compare RULES, ML, and HYBRID decision modes.")
    add_num(doc, "Demonstrate cognitive / adaptive hopping when enabled.")
    add_num(doc, "This is software-only. No radio hardware and no radiated RF are required.")

    # 2
    doc.add_heading("2. Before you start — system requirements", level=1)
    add_num(doc, "A Windows PC with Python installed (Python 3.10 or later recommended).")
    add_num(doc, "Internet access once, to install Python packages the first time.")
    add_num(doc, "A web browser such as Microsoft Edge or Google Chrome.")
    add_num(doc, "The application folder: cognitive_radio_sim")

    # 3
    doc.add_heading("3. How to start the software (first time)", level=1)
    add_num(doc, "Open PowerShell or Command Prompt.")
    add_num(doc, "Change to the application folder by typing:")
    add_p(doc, r"cd C:\work\projectx\embdeeded\proposalprojects\cognitive_radio_sim", italic=True)
    add_num(doc, "Install required packages (first time only) by typing:")
    add_p(doc, "python -m pip install -r requirements.txt", italic=True)
    add_num(doc, "Start the application by typing:")
    add_p(doc, "python -m streamlit run app.py", italic=True)
    add_num(doc, "Wait until the terminal shows a Local URL.")
    add_num(doc, "Open a browser and go to: http://localhost:8501")
    add_num(doc, "If the browser opens automatically, use that window.")
    add_num(doc, "You should see the title: Cognitive Radio Channel Selection under Jamming.")

    # 4
    doc.add_heading("4. How to start the software (later times)", level=1)
    add_num(doc, "Open PowerShell or Command Prompt.")
    add_num(doc, "Go to the application folder (same path as in Section 3).")
    add_num(doc, "Run: python -m streamlit run app.py")
    add_num(doc, "Open http://localhost:8501 in the browser.")

    # 5
    doc.add_heading("5. How to stop the software", level=1)
    add_num(doc, "Click into the terminal window where the app is running.")
    add_num(doc, "Press Ctrl+C to stop the server.")
    add_num(doc, "You may close the browser tab after that.")

    # 6
    doc.add_heading("6. Screen layout — what each area means", level=1)
    add_num(doc, "Left sidebar — configuration, jam injection, and run controls.")
    add_num(doc, "Top metrics — simulation step, active TX channel, TX power, number of channels, jam profile.")
    add_num(doc, "Channel map — one card per channel. Green = FREE, amber = DEGRADED, red = BLOCKED. Gold border = active transmit channel.")
    add_num(doc, "Cognitive Radio response — status code and message.")
    add_num(doc, "Decision detail — engine used, action, recommended channels, power advice, hop set.")
    add_num(doc, "Quick jam toggle — buttons CH0, CH1, … to jam or unjam one channel quickly.")
    add_num(doc, "Recent decision history — last decisions for review.")

    # 7
    doc.add_heading("7. Important rule before changing settings", level=1)
    add_num(doc, "Change the setting in the sidebar (for example N, decision engine, hopping, thresholds, or seed).")
    add_num(doc, "Click Apply configuration.")
    add_num(doc, "Wait if the app says it is training the ML recommender.")
    add_num(doc, "Only then continue with jam injection and Single step / Auto steps.")
    add_p(
        doc,
        "If you skip Apply configuration, the on-screen behaviour may still use the previous settings.",
        italic=True,
    )

    # 8
    doc.add_heading("8. Step-by-step demonstration script", level=1)
    add_p(
        doc,
        "Follow these numbered scenarios in order for a complete demonstration. "
        "After each scenario, check the Expected result before moving to the next.",
    )

    doc.add_heading("8.1 Scenario 1 — Clean spectrum (no jam)", level=2)
    add_num(doc, "In the sidebar, under Jam injection, click Clear jam.")
    add_num(doc, "Under Run control, click Single step once or twice.")
    add_num(doc, "Look at the channel map.")
    add_num(doc, "Expected result: most or all channels are FREE (green). Status is usually OK. Active TX is on a usable channel.")

    doc.add_heading("8.2 Scenario 2 — Spot jam (one channel)", level=2)
    add_num(doc, "Set Jam profile to SPOT.")
    add_num(doc, "In Target channels, select one channel (for example 0).")
    add_num(doc, "Click Apply jam.")
    add_num(doc, "Click Single step.")
    add_num(doc, "Expected result: the selected channel turns BLOCKED (red). Cognitive Radio recommends another free channel. Active TX moves off the jammed channel.")

    doc.add_heading("8.3 Scenario 3 — Multi-spot jam (several channels)", level=2)
    add_num(doc, "Set Jam profile to MULTI_SPOT.")
    add_num(doc, "In Target channels, select two or more channels (for example 0 and 1).")
    add_num(doc, "Click Apply jam.")
    add_num(doc, "Click Single step.")
    add_num(doc, "Expected result: selected channels are BLOCKED. Cognitive Radio recommends a free channel. Message and status code update.")

    doc.add_heading("8.4 Scenario 4 — Sweep jam", level=2)
    add_num(doc, "Set Jam profile to SWEEP.")
    add_num(doc, "Click Apply jam.")
    add_num(doc, "Set Auto steps to 5 or 10.")
    add_num(doc, "Click Run auto steps.")
    add_num(doc, "Watch the channel map while steps advance.")
    add_num(doc, "Expected result: the BLOCKED channel moves over time. Cognitive Radio keeps recommending a usable channel as the jam sweeps.")

    doc.add_heading("8.5 Scenario 5 — Barrage jam and increase power", level=2)
    add_num(doc, "Set Jam profile to BARRAGE.")
    add_num(doc, "Click Apply jam.")
    add_num(doc, "Click Single step several times.")
    add_num(doc, "Watch the Cognitive Radio response and the TX power metric.")
    add_num(doc, "Expected result: many or all channels are BLOCKED. Status becomes INCREASE_POWER if power can still rise. TX power increases step by step.")
    add_num(doc, "If power is already at the maximum, expected status is ALL_BLOCKED.")

    doc.add_heading("8.6 Scenario 6 — Cognitive / adaptive hopping", level=2)
    add_num(doc, "In Configuration, turn Cognitive / adaptive hopping ON.")
    add_num(doc, "Click Apply configuration.")
    add_num(doc, "Set Jam profile to SWEEP or MULTI_SPOT.")
    add_num(doc, "Click Apply jam.")
    add_num(doc, "Run several auto steps.")
    add_num(doc, "Open Decision detail and read Hop set.")
    add_num(doc, "Expected result: hop set is shown; jammed members can be removed and clean channels added. Transmitter follows the hop set.")

    doc.add_heading("8.7 Scenario 7 — Compare RULES, ML, and HYBRID", level=2)
    add_num(doc, "Set Decision engine to RULES.")
    add_num(doc, "Click Apply configuration.")
    add_num(doc, "Apply a MULTI_SPOT jam and click Single step.")
    add_num(doc, "Note the status code, message, and recommended channel.")
    add_num(doc, "Change Decision engine to ML.")
    add_num(doc, "Click Apply configuration, apply the same jam, and step again.")
    add_num(doc, "Change Decision engine to HYBRID and repeat.")
    add_num(doc, "Expected result: the response format stays the same (status, message, channels / power advice). Engine name shows RULES, ML, or HYBRID.")

    doc.add_heading("8.8 Scenario 8 — Quick jam toggle", level=2)
    add_num(doc, "Scroll to Quick jam toggle.")
    add_num(doc, "Click CH0.")
    add_num(doc, "Click Single step if the map does not update immediately.")
    add_num(doc, "Click CH0 again to clear that jam, or click other CH buttons.")
    add_num(doc, "Expected result: individual channels toggle between jammed and not jammed for a live demonstration.")

    doc.add_heading("8.9 Scenario 9 — Change number of channels", level=2)
    add_num(doc, "Set Number of channels (N) to 4, 8, or 16.")
    add_num(doc, "Click Apply configuration.")
    add_num(doc, "Wait for ML training to finish.")
    add_num(doc, "Click Clear jam, then Single step.")
    add_num(doc, "Expected result: the channel map shows exactly N cards.")

    doc.add_heading("8.10 Scenario 10 — Reset to a clean start", level=2)
    add_num(doc, "Click Clear jam.")
    add_num(doc, "Click Reset simulation.")
    add_num(doc, "Wait for retraining if prompted.")
    add_num(doc, "Expected result: simulation returns to a clean starting point for the next reviewer.")

    # 9
    doc.add_heading("9. How to read status codes", level=1)
    add_num(doc, "OK — at least one FREE channel is recommended.")
    add_num(doc, "DEGRADED_OK — no FREE channel; a usable DEGRADED channel is recommended.")
    add_num(doc, "INCREASE_POWER — all channels blocked at current power; raise TX power.")
    add_num(doc, "ALL_BLOCKED — still blocked at maximum allowed power.")
    add_num(doc, "HOLD — keep the current allocation (if shown).")
    add_num(doc, "INVALID_INPUT — observations missing or invalid (should not appear in normal demo use).")

    # 10
    doc.add_heading("10. Sidebar controls reference", level=1)
    add_num(doc, "Number of channels (N) — choose how many logical channels to simulate.")
    add_num(doc, "Decision engine — RULES, ML, or HYBRID.")
    add_num(doc, "Cognitive / adaptive hopping — enable hop-set adaptation.")
    add_num(doc, "Jam power — higher values make blockage stronger.")
    add_num(doc, "SINR FREE threshold — channels at or above this are FREE.")
    add_num(doc, "SINR BLOCKED threshold — channels below this are BLOCKED.")
    add_num(doc, "Maximum TX power — upper limit for the increase-power path.")
    add_num(doc, "Seed — repeatable random sequence for fair comparisons.")
    add_num(doc, "Apply configuration — apply the settings above.")
    add_num(doc, "Jam profile — NONE, SPOT, MULTI_SPOT, SWEEP, or BARRAGE.")
    add_num(doc, "Target channels — channels used by spot / multi-spot profiles.")
    add_num(doc, "Apply jam / Clear jam — inject or remove the selected jam profile.")
    add_num(doc, "Single step — one closed-loop update.")
    add_num(doc, "Run auto steps — several updates in sequence.")
    add_num(doc, "Reset simulation — return to a clean start.")

    # 11
    doc.add_heading("11. Short presenter checklist (one page)", level=1)
    add_num(doc, "Start the app and open http://localhost:8501.")
    add_num(doc, "Show clean spectrum (Clear jam → Single step).")
    add_num(doc, "Jam CH0 and CH1 (MULTI_SPOT) → show red cells → show CR recommendation → show TX move.")
    add_num(doc, "Run SWEEP with auto steps → show blocked cell moving.")
    add_num(doc, "Run BARRAGE → show INCREASE_POWER and rising TX power.")
    add_num(doc, "Optionally enable hopping and show hop-set update.")
    add_num(doc, "Optionally switch RULES → HYBRID and show the same response format.")
    add_num(doc, "Remind the audience: software-only demonstration; no radiated RF.")

    # 12
    doc.add_heading("12. Troubleshooting", level=1)
    add_num(doc, "If streamlit is not recognised, use: python -m streamlit run app.py")
    add_num(doc, "If the page does not load, check the terminal for Local URL and open http://localhost:8501")
    add_num(doc, "If the map does not change after jam, click Single step.")
    add_num(doc, "If settings seem ignored, click Apply configuration after changing them.")
    add_num(doc, "If ML says it is not trained, click Apply configuration or Reset simulation.")
    add_num(doc, "If labels or buttons look hard to read, hard-refresh the browser with Ctrl+F5.")
    add_num(doc, "For a clean restart of the whole demo, Clear jam, then Reset simulation.")

    # 13
    doc.add_heading("13. Closing note", level=1)
    add_p(
        doc,
        "Operate the software using the numbered steps in Section 8 for a consistent demonstration. "
        "Always point reviewers to the status code and message after each action so the Cognitive Radio decision remains inspectable.",
    )

    add_center(doc, "— End of user guide —", 9, False, GRAY)

    OUT_APP.parent.mkdir(parents=True, exist_ok=True)
    OUT_DESK.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_APP)
    doc.save(OUT_DESK)
    print(f"Wrote {OUT_APP}")
    print(f"Wrote {OUT_DESK}")


if __name__ == "__main__":
    main()
