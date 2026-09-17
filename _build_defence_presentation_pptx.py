#!/usr/bin/env python3
"""
Build PowerPoint presentation for Senior Defence Scientists
on the Cognitive Radio Simulation. Includes architecture diagram + live UI screenshot.
Also embeds both images into the Word Presentation Guide.
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches as DocInches, Pt as DocPt, RGBColor as DocRGB

BASE = Path(__file__).resolve().parent
ASSETS = BASE / "assets"
LOGO = ASSETS / "eage_logo_hires.png"
ARCH = ASSETS / "system_architecture_block_diagram.png"
UI = ASSETS / "ui_modular_node_view.png"

OUT_PPTX = BASE / "Cognitive_Radio_Simulation_Presentation_for_Defence_Scientists.pptx"
OUT_PPTX_DESK = Path.home() / "Desktop" / "proposals" / "Cognitive_Radio_Simulation_Presentation_for_Defence_Scientists.pptx"
OUT_DOCX = BASE / "Cognitive_Radio_Simulation_Presentation_Guide_for_Defence_Scientists.docx"
OUT_DOCX_DESK = Path.home() / "Desktop" / "proposals" / "Cognitive_Radio_Simulation_Presentation_Guide_for_Defence_Scientists.docx"

NAVY = RGBColor(0x0B, 0x2A, 0x5B)
SLATE = RGBColor(0x1E, 0x29, 0x3B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xF8, 0xFA, 0xFC)
MUTED = RGBColor(0x64, 0x74, 0x8B)
ACCENT = RGBColor(0x03, 0x69, 0xA1)
GREEN = RGBColor(0x15, 0x80, 0x3D)
RED = RGBColor(0xB9, 0x1C, 0x1C)


def set_run_font(run, size=18, bold=False, color=SLATE):
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_bg(slide, color=NAVY):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    # send to back
    spTree = slide.shapes._spTree
    sp = shape._element
    spTree.remove(sp)
    spTree.insert(2, sp)


def add_top_bar(slide):
    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.08))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()


def add_footer(slide, page, total=12):
    box = slide.shapes.add_textbox(Inches(0.5), Inches(7.15), Inches(10), Inches(0.3))
    tf = box.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = "eAge Innovations  |  Cognitive Radio Simulation  |  Defence Scientific Briefing"
    set_run_font(run, size=10, color=MUTED)
    right = slide.shapes.add_textbox(Inches(11.5), Inches(7.15), Inches(1.4), Inches(0.3))
    tf2 = right.text_frame
    p2 = tf2.paragraphs[0]
    p2.alignment = PP_ALIGN.RIGHT
    run2 = p2.add_run()
    run2.text = f"{page} / {total}"
    set_run_font(run2, size=10, color=MUTED)


def add_title(slide, text, top=0.25, size=28):
    box = slide.shapes.add_textbox(Inches(0.5), Inches(top), Inches(12.3), Inches(0.6))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    set_run_font(run, size=size, bold=True, color=NAVY)
    return box


def add_subtitle(slide, text, top=0.75):
    box = slide.shapes.add_textbox(Inches(0.5), Inches(top), Inches(12.3), Inches(0.4))
    tf = box.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    set_run_font(run, size=14, color=MUTED)
    return box


def add_numbered_body(slide, items, left=0.5, top=1.3, width=12.3, height=5.5, size=16):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        p.level = 0
        run = p.add_run()
        run.text = item
        set_run_font(run, size=size, bold=False, color=SLATE)
    return box


def blank_slide(prs):
    layout = prs.slide_layouts[6]  # blank
    return prs.slides.add_slide(layout)


def build_pptx():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    total = 12

    # ---------- SLIDE 1: Title ----------
    s = blank_slide(prs)
    add_bg(s, NAVY)
    if LOGO.exists():
        # white card behind logo
        card = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(10.4), Inches(0.35), Inches(2.5), Inches(0.85))
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.fill.background()
        s.shapes.add_picture(str(LOGO), Inches(10.65), Inches(0.48), height=Inches(0.55))
    box = s.shapes.add_textbox(Inches(0.7), Inches(2.2), Inches(11.5), Inches(1.2))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = "COGNITIVE RADIO CHANNEL SELECTION\nUNDER HOSTILE JAMMING"
    set_run_font(run, size=32, bold=True, color=WHITE)
    box2 = s.shapes.add_textbox(Inches(0.7), Inches(3.7), Inches(11.5), Inches(1.0))
    tf2 = box2.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    r2 = p2.add_run()
    r2.text = "Software Demonstration for Senior Defence Scientists\nSense · Decide · Act — Interference Capture · Hybrid ML Safety · Cognitive Adaptive Frequency Hopping"
    set_run_font(r2, size=16, color=RGBColor(0xCB, 0xD5, 0xE1))
    box3 = s.shapes.add_textbox(Inches(0.7), Inches(5.5), Inches(11.5), Inches(0.8))
    p3 = box3.text_frame.paragraphs[0]
    r3 = p3.add_run()
    r3.text = "eAge Innovations  |  Defence SDR & Electronic Warfare Lab  |  September 2026\nPrepared for Shri S. K. Sir & Senior Defence Scientists"
    set_run_font(r3, size=13, color=RGBColor(0x94, 0xA3, 0xB8))

    # ---------- SLIDE 2: Agenda ----------
    s = blank_slide(prs)
    add_top_bar(s)
    add_title(s, "1. Presentation Agenda")
    add_subtitle(s, "Numbered flow for the live scientific briefing")
    add_numbered_body(
        s,
        [
            "1.1  Objective and SK priorities for this demonstration",
            "1.2  Closed-loop system architecture (block diagram)",
            "1.3  Live modular UI — Jammer, Receiver, Transmitter + ML",
            "1.4  How interference is captured (ES sensing)",
            "1.5  Hybrid ML decision with deterministic safety gate",
            "1.6  Dynamic frequency hopping based on cognitive decisions",
            "1.7  Recommended live demo sequence and closing",
        ],
        top=1.4,
        size=18,
    )
    add_footer(s, 2, total)

    # ---------- SLIDE 3: Objective ----------
    s = blank_slide(prs)
    add_top_bar(s)
    add_title(s, "2. What This Demonstration Proves")
    add_subtitle(s, "Aligned to the look-and-feel and technical priorities discussed with Shri S. K. Sir")
    add_numbered_body(
        s,
        [
            "2.1  Present the agreed modular look-and-feel of the Cognitive Radio closed loop.",
            "2.2  Show capturing interference through Electronic Support (ES) sensing and FREE / DEGRADED / BLOCKED classification.",
            "2.3  Show cognitive decisions using RULES / ML / HYBRID, with HYBRID preferred for mission assurance.",
            "2.4  Show dynamic frequency hopping that updates the hop pool from live cognitive decisions.",
            "2.5  Keep every claim tied to observable screen evidence and an explainability trace.",
            "2.6  Position this as a software decision twin ready for later eADM / SDR integration — not a flight-certified article today.",
        ],
        top=1.35,
        size=17,
    )
    add_footer(s, 3, total)

    # ---------- SLIDE 4: Architecture image ----------
    s = blank_slide(prs)
    add_top_bar(s)
    add_title(s, "3. Closed-Loop System Architecture", top=0.18, size=24)
    add_subtitle(s, "Figure 1 — Sense–Decide–Act under Electronic Attack", top=0.65)
    if ARCH.exists():
        # wide image centered
        s.shapes.add_picture(str(ARCH), Inches(0.35), Inches(1.1), width=Inches(12.6))
    add_footer(s, 4, total)

    # ---------- SLIDE 5: Architecture numbered explanation ----------
    s = blank_slide(prs)
    add_top_bar(s)
    add_title(s, "4. Architecture — Numbered Explanation")
    add_subtitle(s, "How to narrate Figure 1 to the panel")
    add_numbered_body(
        s,
        [
            "4.1  Jammer (Red): adversary ECM — Spot, Multi-Spot, Sweep, Barrage.",
            "4.2  Contested RF Medium: friendly signal + jammer power + noise/fading.",
            "4.3  Receiver / ES Node (Blue): measures SINR and packages CRRequest telemetry.",
            "4.4  Cognitive Decision Engine & ML (Slate): 64×32 MLP proposal + Hybrid Safety Gate + adaptive hop synthesis.",
            "4.5  Transmitter / ECCM Actuator (Green): retunes carrier, regulates power, executes hop pool.",
            "4.6  Closed loop: radiated friendly carrier returns to the contested medium.",
            "4.7  Key message: Sense → Decide → Act, with explainable safety containment of ML.",
        ],
        top=1.35,
        size=17,
    )
    add_footer(s, 5, total)

    # ---------- SLIDE 6: Live UI screenshot ----------
    s = blank_slide(prs)
    add_top_bar(s)
    add_title(s, "5. Live Software UI — Modular Node View", top=0.18, size=24)
    add_subtitle(s, "Figure 2 — Jammer · Receiver · Transmitter with Attached ML Policy", top=0.65)
    if UI.exists():
        s.shapes.add_picture(str(UI), Inches(0.45), Inches(1.05), width=Inches(12.4))
    add_footer(s, 6, total)

    # ---------- SLIDE 7: How to read the UI ----------
    s = blank_slide(prs)
    add_top_bar(s)
    add_title(s, "6. How to Read the Live UI")
    add_subtitle(s, "Use this while pointing at Figure 2 during the demonstration")
    # three columns as cards
    cards = [
        (0.4, RED, "6.1  JAMMER (Red)", [
            "ECM threat profile control",
            "Jam power and emit / silence",
            "Quick presets: Spot / Multi",
            "Shows adversary activity state",
        ]),
        (4.55, ACCENT, "6.2  RECEIVER (Blue)", [
            "FREE / DEGRADED / BLOCKED counts",
            "Channel Spectrum State Map",
            "Cleanest vs worst impaired link",
            "This is interference capture",
        ]),
        (8.7, GREEN, "6.3  TRANSMITTER (Green)", [
            "Active carrier and TX power",
            "Hopping mode / hop pool",
            "Attached MLP recommender",
            "Proposal, confidence, action",
        ]),
    ]
    for left, color, title, bullets in cards:
        shape = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(left), Inches(1.35), Inches(3.95), Inches(5.2))
        shape.fill.solid()
        shape.fill.fore_color.rgb = LIGHT
        shape.line.color.rgb = color
        shape.line.width = Pt(2)
        hdr = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(left), Inches(1.35), Inches(3.95), Inches(0.55))
        hdr.fill.solid()
        hdr.fill.fore_color.rgb = color
        hdr.line.fill.background()
        tb = s.shapes.add_textbox(Inches(left + 0.15), Inches(1.42), Inches(3.65), Inches(0.4))
        run = tb.text_frame.paragraphs[0].add_run()
        run.text = title
        set_run_font(run, size=14, bold=True, color=WHITE)
        body = s.shapes.add_textbox(Inches(left + 0.2), Inches(2.1), Inches(3.55), Inches(4.2))
        tf = body.text_frame
        tf.word_wrap = True
        for i, b in enumerate(bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(12)
            r = p.add_run()
            r.text = f"{i + 1}.  {b}"
            set_run_font(r, size=14, color=SLATE)
    add_footer(s, 7, total)

    # ---------- SLIDE 8: Interference + FH ----------
    s = blank_slide(prs)
    add_top_bar(s)
    add_title(s, "7. Interference Capture & Cognitive Dynamic FH")
    add_subtitle(s, "The two technical capabilities emphasised by Shri S. K. Sir")
    add_numbered_body(
        s,
        [
            "7.1  Capturing interference: Receiver measures jam power and SINR per channel.",
            "7.2  Classification: FREE (≥10 dB), DEGRADED (3–10 dB), BLOCKED (<3 dB).",
            "7.3  Decisions use measured evidence only — no oracle jammer identity.",
            "7.4  Dynamic FH: when hopping is enabled, adapt_hop_set rebuilds the hop pool.",
            "7.5  Threatened (BLOCKED) hop members are evicted; clean usable channels are filled in.",
            "7.6  Result: hopping follows cognitive decisions, not a blind fixed PRFH list.",
        ],
        top=1.35,
        size=17,
    )
    add_footer(s, 8, total)

    # ---------- SLIDE 9: ML ----------
    s = blank_slide(prs)
    add_top_bar(s)
    add_title(s, "8. Machine Learning in This Simulation")
    add_subtitle(s, "Advisory recommender inside a hybrid safety architecture")
    add_numbered_body(
        s,
        [
            "8.1  Model: supervised MLP classifier (64 × 32), ReLU, Adam.",
            "8.2  Input: 3N+2 feature vector (for N=8 → 26 features: SINR, state, jam, power, headroom).",
            "8.3  Output: recommended channel or INCREASE_POWER, with Softmax confidence.",
            "8.4  Modes: RULES · ML · HYBRID (preferred for defence review).",
            "8.5  Safety Gate: veto BLOCKED proposals and out-of-headroom power increases.",
            "8.6  Assurance message: ML advises; deterministic rules contain; every override is logged.",
        ],
        top=1.35,
        size=17,
    )
    add_footer(s, 9, total)

    # ---------- SLIDE 10: Demo sequence ----------
    s = blank_slide(prs)
    add_top_bar(s)
    add_title(s, "9. Recommended Live Demo Sequence")
    add_subtitle(s, "Numbered order for a 20–25 minute presentation")
    add_numbered_body(
        s,
        [
            "9.1  Opening: state Sense–Decide–Act and the three proofs.",
            "9.2  Walk the architecture diagram (Figure 1), then show the live UI (Figure 2).",
            "9.3  Benign baseline — all FREE, status OK.",
            "9.4  Spot jam on active carrier — show BLOCKED capture and HYBRID alternate channel.",
            "9.5  Enable cognitive hopping — jam a hop-set member — show eviction and refill.",
            "9.6  Barrage — show INCREASE_POWER within headroom.",
            "9.7  Close with explainability Stages 2–3 and invite questions.",
        ],
        top=1.35,
        size=17,
    )
    add_footer(s, 10, total)

    # ---------- SLIDE 11: Phrases ----------
    s = blank_slide(prs)
    add_top_bar(s)
    add_title(s, "10. Preferred Speaking Points")
    add_numbered_body(
        s,
        [
            "10.1  Sense–Decide–Act closed loop under Electronic Attack.",
            "10.2  Interference capture through Electronic Support sensing.",
            "10.3  Hybrid safety-gated machine learning recommender.",
            "10.4  Cognitive adaptive frequency hopping / dynamic hop-pool management.",
            "10.5  Mission assurance with a human-readable audit trail.",
            "10.6  Software decision twin prepared for later eADM / SDR integration.",
            "10.7  Avoid claiming airborne certification or absolute jam-proof performance.",
        ],
        top=1.3,
        size=17,
    )
    add_footer(s, 11, total)

    # ---------- SLIDE 12: Close ----------
    s = blank_slide(prs)
    add_bg(s, NAVY)
    box = s.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.7), Inches(1.2))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "Closing Summary for the Panel"
    set_run_font(run, size=30, bold=True, color=WHITE)
    box2 = s.shapes.add_textbox(Inches(0.8), Inches(3.3), Inches(11.7), Inches(2.5))
    tf = box2.text_frame
    tf.word_wrap = True
    lines = [
        "1.  Agreed look-and-feel is presented in modular architecture and live UI.",
        "2.  Interference is captured by ES sensing and channel-state classification.",
        "3.  Decisions use hybrid ML under a deterministic safety gate.",
        "4.  Frequency hopping updates dynamically from cognitive decisions.",
        "5.  We request guidance on priorities for the next eADM laboratory step.",
    ]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        r = p.add_run()
        r.text = line
        set_run_font(r, size=18, color=RGBColor(0xE2, 0xE8, 0xF0))
    box3 = s.shapes.add_textbox(Inches(0.8), Inches(6.3), Inches(11.7), Inches(0.5))
    r3 = box3.text_frame.paragraphs[0].add_run()
    r3.text = "Thank you, Sir. We are ready for questions."
    set_run_font(r3, size=16, bold=True, color=RGBColor(0x93, 0xC5, 0xFD))

    prs.save(str(OUT_PPTX))
    print(f"Saved PPTX: {OUT_PPTX}")
    if OUT_PPTX_DESK.parent.exists():
        prs.save(str(OUT_PPTX_DESK))
        print(f"Copied PPTX: {OUT_PPTX_DESK}")


def set_doc_run(run, size=10.5, bold=False, italic=False, color=DocRGB(0x1F, 0x29, 0x37)):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = DocPt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color


def add_doc_p(doc, text="", size=10.5, bold=False, italic=False, color=DocRGB(0x1F, 0x29, 0x37), space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = DocPt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if text:
        r = p.add_run(text)
        set_doc_run(r, size=size, bold=bold, italic=italic, color=color)
    return p


def add_doc_h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = DocPt(14)
    p.paragraph_format.space_after = DocPt(6)
    r = p.add_run(text)
    set_doc_run(r, size=13, bold=True, color=DocRGB(0x0B, 0x2A, 0x5B))
    return p


def add_doc_num(doc, num, title, body):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = DocInches(0.32)
    p.paragraph_format.first_line_indent = DocInches(-0.32)
    p.paragraph_format.space_before = DocPt(2)
    p.paragraph_format.space_after = DocPt(5)
    p.paragraph_format.line_spacing = 1.15
    r_n = p.add_run(f"{num} ")
    set_doc_run(r_n, size=10.5, bold=True, color=DocRGB(0x0B, 0x2A, 0x5B))
    if title:
        r_t = p.add_run(f"{title} ")
        set_doc_run(r_t, size=10.5, bold=True, color=DocRGB(0x1E, 0x29, 0x3B))
    r_b = p.add_run(body)
    set_doc_run(r_b, size=10.5, bold=False, color=DocRGB(0x1F, 0x29, 0x37))
    return p


def add_doc_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = 1  # center
    p.paragraph_format.space_before = DocPt(4)
    p.paragraph_format.space_after = DocPt(10)
    r = p.add_run(text)
    set_doc_run(r, size=9, bold=True, italic=True, color=DocRGB(0x64, 0x74, 0x8B))
    return p


def build_docx_with_images():
    """Rebuild presentation guide Word doc with both figures embedded."""
    doc = Document()
    for s in doc.sections:
        s.top_margin = DocInches(0.75)
        s.bottom_margin = DocInches(0.75)
        s.left_margin = DocInches(0.8)
        s.right_margin = DocInches(0.8)

    add_doc_p(doc, "eAge Innovations — Defence SDR & Electronic Warfare Lab", size=10, bold=True, color=DocRGB(0x0B, 0x2A, 0x5B), space_after=2)
    add_doc_p(
        doc,
        "Reference: EAGE-CR-PRES-2026-001  |  Audience: Shri S. K. Sir & Senior Defence Scientists  |  September 2026",
        size=8.5,
        color=DocRGB(0x64, 0x74, 0x8B),
        space_after=8,
    )
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_after = DocPt(10)
    p_div._element.get_or_add_pPr().append(
        parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="18" w:space="1" w:color="0B2A5B"/></w:pBdr>')
    )

    t = doc.add_paragraph()
    t.paragraph_format.space_after = DocPt(3)
    set_doc_run(t.add_run("HOW TO PRESENT THE COGNITIVE RADIO SIMULATION"), size=15, bold=True, color=DocRGB(0x0B, 0x2A, 0x5B))

    add_doc_num(
        doc,
        "0.1",
        "Document format:",
        "This guide uses numbered points only. Figures below are used in both this Word document and the companion PowerPoint.",
    )
    add_doc_num(
        doc,
        "0.2",
        "Companion PowerPoint:",
        "Cognitive_Radio_Simulation_Presentation_for_Defence_Scientists.pptx",
    )

    add_doc_h1(doc, "1. Objective of This Presentation")
    add_doc_num(doc, "1.1", "Primary goal:", "Show Sense–Decide–Act under jamming with agreed look-and-feel, interference capture, and cognitive dynamic FH.")
    add_doc_num(doc, "1.2", "Desired outcome:", "Panel understands the software decision twin and path to eADM / SDR integration.")
    add_doc_num(doc, "1.3", "Scope limit:", "Not a claim of airborne certification or full field RF trial.")
    add_doc_num(doc, "1.4", "SK priorities:", "Look-and-feel; capturing interference; dynamic FH based on cognitive decisions.")

    add_doc_h1(doc, "2. Figure 1 — Closed-Loop System Architecture")
    add_doc_num(doc, "2.1", "Purpose of figure:", "Use this diagram first to explain the closed loop before opening the live software.")
    if ARCH.exists():
        p = doc.add_paragraph()
        p.alignment = 1
        p.add_run().add_picture(str(ARCH), width=DocInches(6.6))
        add_doc_caption(doc, "Figure 1: Closed-Loop Cognitive Radio System Architecture under Electronic Attack.")
    add_doc_num(doc, "2.2", "Narration point 1:", "Jammer injects hostile ECM power into the contested RF medium.")
    add_doc_num(doc, "2.3", "Narration point 2:", "Receiver performs ES sensing and issues CRRequest telemetry.")
    add_doc_num(doc, "2.4", "Narration point 3:", "Cognitive Decision Engine & ML apply hybrid safety and adaptive hop logic.")
    add_doc_num(doc, "2.5", "Narration point 4:", "Transmitter actuates approved carrier / power / hop set and closes the RF loop.")

    add_doc_h1(doc, "3. Figure 2 — Live Modular Node UI")
    add_doc_num(doc, "3.1", "Purpose of figure:", "Use this screenshot to show the agreed look-and-feel of Jammer, Receiver, and Transmitter with attached ML.")
    if UI.exists():
        p = doc.add_paragraph()
        p.alignment = 1
        p.add_run().add_picture(str(UI), width=DocInches(6.6))
        add_doc_caption(doc, "Figure 2: Live Cognitive Radio Simulation — Modular System Architecture & Node View.")
    add_doc_num(doc, "3.2", "Left panel — Jammer:", "ECM threat profile, jam power, emit/silence, quick Spot / Multi presets.")
    add_doc_num(doc, "3.3", "Centre panel — Receiver:", "FREE / DEGRADED / BLOCKED counts and Channel Spectrum State Map — this is interference capture.")
    add_doc_num(doc, "3.4", "Right panel — Transmitter:", "Active carrier, power, hopping mode, and Attached MLP Policy with proposal, confidence, and dispatched action.")
    add_doc_num(doc, "3.5", "Demo tip:", "While speaking, point physically to each coloured column so the panel links architecture (Figure 1) to live UI (Figure 2).")

    add_doc_h1(doc, "4. Before the Meeting — Preparation Checklist")
    add_doc_num(doc, "4.1", "Start the application:", "python -m streamlit run app.py --server.port 8501")
    add_doc_num(doc, "4.2", "Open browser:", "http://localhost:8501")
    add_doc_num(doc, "4.3", "Preferred view:", "System Architecture & Node View")
    add_doc_num(doc, "4.4", "Defaults:", "N=8; Decision engine=HYBRID; Hopping OFF initially; Light theme if bright room")
    add_doc_num(doc, "4.5", "Train ML once:", "Complete training before audience arrival")
    add_doc_num(doc, "4.6", "Reset jammer:", "NONE / clear for benign opening")
    add_doc_num(doc, "4.7", "Keep ready:", "This Word guide, the PowerPoint, Scenarios Brief, ML Q&A Brief")
    add_doc_num(doc, "4.8", "Projector:", "Browser zoom 110–125% if needed")
    add_doc_num(doc, "4.9", "Time budget:", "20–25 minutes demo + 10–15 minutes questions")

    add_doc_h1(doc, "5. Opening Remarks")
    add_doc_num(doc, "5.1", "Greeting:", "Sir, today we demonstrate Cognitive Radio software simulation under contested EW conditions.")
    add_doc_num(doc, "5.2", "Positioning:", "PC software closed-loop twin; later eADM / SDR can replace the channel simulator using the same decision interfaces.")
    add_doc_num(doc, "5.3", "Three proofs:", "Interference capture; hybrid ML safety decisions; dynamic FH from cognitive decisions.")
    add_doc_num(doc, "5.4", "Invite interaction:", "Please interrupt to change jammer profile or decision mode.")

    add_doc_h1(doc, "6. Live Demonstration Sequence")
    add_doc_num(doc, "6.1", "Action:", "Show Figure 1 architecture, then switch to live UI (Figure 2).")
    add_doc_num(doc, "6.2", "Action:", "Benign baseline — jammer NONE — status OK.")
    add_doc_num(doc, "6.3", "Action:", "Spot jam active carrier — show BLOCKED capture and HYBRID alternate channel.")
    add_doc_num(doc, "6.4", "Action:", "Enable cognitive hopping — jam hop-set member — show eviction and refill.")
    add_doc_num(doc, "6.5", "Action:", "Barrage — show INCREASE_POWER within headroom.")
    add_doc_num(doc, "6.6", "Action:", "Explain Stages 2–3 of the Decision Engine trace (ML proposal + safety).")

    add_doc_h1(doc, "7. Machine Learning Speaking Points")
    add_doc_num(doc, "7.1", "Model:", "Supervised MLP 64×32, ReLU, Adam.")
    add_doc_num(doc, "7.2", "Input:", "3N+2 features (N=8 → 26 normalised sensing features).")
    add_doc_num(doc, "7.3", "Output:", "Channel recommendation or INCREASE_POWER with Softmax confidence.")
    add_doc_num(doc, "7.4", "Preferred mode:", "HYBRID — ML proposes, deterministic safety validates.")
    add_doc_num(doc, "7.5", "Assurance:", "ML is advisory; safety gate prevents hops into BLOCKED channels.")

    add_doc_h1(doc, "8. Closing Statement")
    add_doc_num(
        doc,
        "8.1",
        "Closing:",
        "Sir, the simulation presents the agreed look-and-feel; captures interference through ES sensing; decides with hybrid ML under safety; and performs dynamic FH based on cognitive decisions. We request guidance for the next eADM laboratory step.",
    )
    add_doc_num(doc, "8.2", "End:", "Thank you, Sir. We are ready for questions.")

    add_doc_p(doc, "")
    add_doc_p(doc, "Prepared by: eAge Innovations Technical Engineering Team", size=9, bold=True, color=DocRGB(0x0B, 0x2A, 0x5B))
    add_doc_p(doc, "For presentation to: Shri S. K. Sir & Senior Defence Scientists", size=9)

    doc.save(str(OUT_DOCX))
    print(f"Saved DOCX: {OUT_DOCX}")
    if OUT_DOCX_DESK.parent.exists():
        doc.save(str(OUT_DOCX_DESK))
        print(f"Copied DOCX: {OUT_DOCX_DESK}")


if __name__ == "__main__":
    build_pptx()
    build_docx_with_images()
