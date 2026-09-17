#!/usr/bin/env python3
"""Build Word (and optional PDF) for ML Fundamentals Q&A mapped to CR simulation code."""

from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

BASE = Path(__file__).resolve().parent
OUT_DOCX = BASE / "Cognitive_Radio_Simulation_ML_Fundamentals_QA.docx"
OUT_DESK = Path.home() / "Desktop" / "proposals" / "Cognitive_Radio_Simulation_ML_Fundamentals_QA.docx"

NAVY = RGBColor(0x0B, 0x2A, 0x5B)
SLATE = RGBColor(0x1E, 0x29, 0x3B)
BODY = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x64, 0x74, 0x8B)

QA_ITEMS = [
    (
        "Data first — Where does the data for this model come from, and is it enough, clean, and representative?",
        "In this simulation, training data is generated synthetically inside the same channel / ECM simulator "
        "(ChannelRecommenderML.train_synthetic in ml_models.py). About 800–1200 labelled cases are created per "
        "training run, covering jam profiles NONE, SPOT, MULTI_SPOT, SWEEP, and BARRAGE, with varied transmit power "
        "and active carrier. Features are numeric and cleaned by construction (no missing RF packets). The data is "
        "representative of the software contested-spectrum scenarios in this demo; it is not yet field-recorded I/Q "
        "from eADM hardware. For the PC demonstration this is sufficient; for deployed EW learning, later work would "
        "add real sensed telemetry while keeping the same feature interface.",
    ),
    (
        "Problem type — Is this classification, regression, clustering, or generation?",
        "This is a supervised multi-class classification problem. The model predicts one discrete action class: "
        "recommended channel id 0 … N−1, or a special class −1 meaning INCREASE_POWER. It does not regress a "
        "continuous SINR estimate, does not cluster threats, and does not generate waveforms. This choice matches "
        "the Cognitive Radio actuation need: pick a channel or advise a power increase.",
    ),
    (
        "Features — What input variables does the model learn from, and why do they matter?",
        "The feature vector length is 3N + 2. For each of N channels the model receives SINR (dB), a state encoding "
        "(FREE=2, DEGRADED=1, BLOCKED=0), and measured jam power; it also receives current transmit power and "
        "remaining power headroom (P_max − P_current). For the default N = 8 this is 26 features, matching the UI "
        "“Input Vector: 26 normalized sensing features.” Features are standardised with StandardScaler before MLP "
        "inference. Feature quality matters more than model size here: if sensing states and jam power are wrong, "
        "the recommender cannot propose safe channels.",
    ),
    (
        "Training vs testing — How is data split so the model is checked on unseen cases?",
        "Today the demonstrator trains on the full synthetic set and reports training accuracy via sklearn "
        "model.score(X_train, y_train) after fit. A formal held-out test split is not yet enforced in "
        "train_synthetic (honest limitation for scientific review). Generalisation is checked operationally in the "
        "live UI on new jam scenarios the operator creates (Spot, Multi, Sweep, Barrage) that were not a single "
        "frozen test file. Recommended next hardening: split synthetic data into train/validation/test "
        "(for example 70/15/15) and report test accuracy / confusion matrix beside train accuracy.",
    ),
    (
        "Algorithm / model — What mathematical structure learns the patterns?",
        "The model is a feed-forward Multi-Layer Perceptron (MLP) classifier from scikit-learn: "
        "MLPClassifier(hidden_layer_sizes=(64, 32), activation=\"relu\", solver=\"adam\", max_iter=400, "
        "random_state=seed). Two hidden layers (64 and 32 neurons) map the feature vector to class logits over "
        "channels plus power-increase. This is intentionally a small, inspectable network suitable for a transparent "
        "defence software demo, not a large deep spectrogram model.",
    ),
    (
        "Loss function — How does the model measure how wrong its predictions are?",
        "For this multi-class MLP classifier, training minimises a cross-entropy / log-loss style objective over the "
        "predicted class probabilities (scikit-learn’s MLP classification loss). In plain terms: if the true label is "
        "“use CH2” and the model assigns low probability to CH2, the loss is high; training pushes probability mass "
        "toward the correct class. Softmax probabilities shown in the UI (and Stage 2 confidence) come from this "
        "probabilistic classification head.",
    ),
    (
        "Optimization — How does the model adjust its internal weights?",
        "Weights are updated iteratively by the Adam optimiser (solver=\"adam\"), which is an adaptive "
        "gradient-based method. Each training iteration reduces the classification loss on the synthetic samples "
        "until convergence or max_iter is reached. A fixed random seed makes training repeatable for demos. After "
        "training, inference is a single forward pass: scale features → MLP → class prediction + probabilities.",
    ),
    (
        "Evaluation metrics — How do we judge whether the model is good?",
        "The UI / training return currently emphasises classification accuracy on the training synthetic set "
        "(train_accuracy). Live demo evaluation is qualitative and scenario-based: correct avoidance of BLOCKED "
        "channels, sensible alternate FREE channel under Spot/Multi jam, INCREASE_POWER under Barrage when headroom "
        "remains, and safety overrides when ML proposes an unsafe channel. For a stronger scientific scorecard, add "
        "precision/recall/F1 per class, confusion matrix, and held-out test accuracy; for EW ops, also track mission "
        "metrics such as fraction of steps on FREE channels and time-to-escape after jam onset.",
    ),
    (
        "Overfitting vs underfitting — How are these risks handled in this design?",
        "Underfitting risk is reduced by using a non-linear MLP (64×32) rather than a single linear rule, and by "
        "training across diverse jam profiles and powers. Overfitting risk exists because current scoring is mainly "
        "on training data and the labeler is itself a rule-based expert function—so the MLP can largely emulate that "
        "expert. Mitigation in this architecture is mission-critical: the Hybrid Safety Gate vetoes BLOCKED or "
        "otherwise unsafe proposals even if the neural net is overconfident. Also, features are low-dimensional and "
        "standardised, and the network is small. Stronger ML hygiene later: held-out test set, early stopping, and "
        "comparing RULES vs ML vs HYBRID on fixed scenario suites.",
    ),
    (
        "Deployment & monitoring — How is the trained model used in operation, and what would monitoring mean here?",
        "In this project, deployment is the Streamlit Cognitive Radio simulation: after startup/reset/apply-config "
        "training, the MLP is used each simulation step inside CognitiveRadioEngine under ML or HYBRID mode, with "
        "proposals shown in the Transmitter Attached ML Policy card and Stage 2 of the explainability trace. "
        "Continuous monitoring in the demo sense is the live FREE/DEGRADED/BLOCKED map, confidence, dispatched "
        "action, and safety-override messages. For a later eADM / SDR deployment, monitoring would track data drift "
        "(new ECM behaviours), confidence calibration, override rate, and link outcomes, then trigger retraining "
        "when sensed distributions diverge from the synthetic training regime.",
    ),
]


def set_run(run, size=10.5, bold=False, italic=False, color=BODY):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color


def add_p(doc, text="", size=10.5, bold=False, italic=False, color=BODY, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if text:
        r = p.add_run(text)
        set_run(r, size=size, bold=bold, italic=italic, color=color)
    return p


def add_qa(doc, num, question, answer):
    # Question with number
    pq = doc.add_paragraph()
    pq.paragraph_format.left_indent = Inches(0.28)
    pq.paragraph_format.first_line_indent = Inches(-0.28)
    pq.paragraph_format.space_before = Pt(8)
    pq.paragraph_format.space_after = Pt(2)
    pq.paragraph_format.line_spacing = 1.15
    rq_n = pq.add_run(f"{num}. ")
    set_run(rq_n, size=10.5, bold=True, color=NAVY)
    rq_t = pq.add_run("Question: ")
    set_run(rq_t, size=10.5, bold=True, color=SLATE)
    rq_b = pq.add_run(question)
    set_run(rq_b, size=10.5, bold=False, color=BODY)

    # Answer with same question number
    pa = doc.add_paragraph()
    pa.paragraph_format.left_indent = Inches(0.28)
    pa.paragraph_format.first_line_indent = Inches(-0.28)
    pa.paragraph_format.space_before = Pt(0)
    pa.paragraph_format.space_after = Pt(6)
    pa.paragraph_format.line_spacing = 1.15
    ra_n = pa.add_run(f"{num}. ")
    set_run(ra_n, size=10.5, bold=True, color=NAVY)
    ra_t = pa.add_run("Answer: ")
    set_run(ra_t, size=10.5, bold=True, color=SLATE)
    ra_b = pa.add_run(answer)
    set_run(ra_b, size=10.5, bold=False, color=BODY)


def build():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.75)
        s.bottom_margin = Inches(0.75)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    add_p(doc, "eAge Innovations — Defence SDR & Electronic Warfare Lab", size=10, bold=True, color=NAVY, space_after=2)
    add_p(
        doc,
        "Reference: EAGE-CR-ML-FUND-QA-2026-001  |  Audience: Senior Defence Scientists  |  September 2026",
        size=8.5,
        color=MUTED,
        space_after=8,
    )
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_after = Pt(10)
    p_div._element.get_or_add_pPr().append(
        parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="18" w:space="1" w:color="0B2A5B"/></w:pBdr>')
    )

    t = doc.add_paragraph()
    t.paragraph_format.space_after = Pt(3)
    set_run(t.add_run("COGNITIVE RADIO SIMULATION — MACHINE LEARNING FUNDAMENTALS Q&A"), size=14, bold=True, color=NAVY)

    s = doc.add_paragraph()
    s.paragraph_format.space_after = Pt(8)
    set_run(
        s.add_run(
            "Ten core ML concepts mapped to what is implemented in this Cognitive Radio software simulation. "
            "Each question is numbered; each answer uses the same question number."
        ),
        size=10.5,
        bold=True,
        color=SLATE,
    )

    add_p(doc, "Primary code references: cr_sim/ml_models.py, cr_sim/cr_engine.py, cr_sim/channel.py, cr_sim/orchestrator.py, app.py", size=9, italic=True, color=MUTED, space_after=10)

    for i, (q, a) in enumerate(QA_ITEMS, start=1):
        add_qa(doc, i, q, a)

    add_p(doc, "")
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(8)
    set_run(h.add_run("Closing note for scientists"), size=11.5, bold=True, color=NAVY)

    add_p(
        doc,
        "1. This codebase implements an explainable ML loop for channel / power classification on synthetic contested-spectrum features.",
        space_after=3,
    )
    add_p(
        doc,
        "2. The main scientific gap to call out honestly is formal train/test split reporting, which should be added before claiming strong generalisation metrics.",
        space_after=3,
    )
    add_p(
        doc,
        "3. Defence assurance in this design does not rely on ML alone: hybrid safety containment remains mandatory around the neural proposal.",
        space_after=10,
    )

    add_p(doc, "Prepared by: eAge Innovations Technical Engineering Team", size=9, bold=True, color=NAVY)
    add_p(doc, "For scientific evaluation by: Shri S. K. Sir & Senior Defence Scientists", size=9, color=BODY)

    doc.save(str(OUT_DOCX))
    print(f"Saved DOCX: {OUT_DOCX}")
    if OUT_DESK.parent.exists():
        try:
            doc.save(str(OUT_DESK))
            print(f"Copied DOCX: {OUT_DESK}")
        except PermissionError:
            alt = OUT_DESK.with_name(OUT_DESK.stem + "_v2.docx")
            doc.save(str(alt))
            print(f"Desktop file locked; saved: {alt}")

    # Optional PDF via docx2pdf (Windows Word) or reportlab fallback skip
    pdf_path = OUT_DOCX.with_suffix(".pdf")
    desk_pdf = OUT_DESK.with_suffix(".pdf")
    try:
        from docx2pdf import convert

        convert(str(OUT_DOCX), str(pdf_path))
        print(f"Saved PDF: {pdf_path}")
        if desk_pdf.parent.exists():
            try:
                convert(str(OUT_DOCX), str(desk_pdf))
                print(f"Copied PDF: {desk_pdf}")
            except Exception as e:
                print(f"Desktop PDF skip: {e}")
    except Exception as e:
        print(f"PDF conversion not available ({e}). Word document is ready.")


if __name__ == "__main__":
    build()
